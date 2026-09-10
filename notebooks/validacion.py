"""
Motor de validación de fuentes + reconstrucción parcial de la matriz + concordancia.

Se apoya en config/fuentes.yaml (mapa de fuentes verificado contra los datos reales
de julio 2026) y compara lo reconstruido contra la matriz objetivo ya diligenciada
(docs/Objetivo/...xlsm).

Uso rápido:
    python notebooks/validacion.py            # corre todo sobre julio y escribe el reporte
    # o desde el notebook: import validacion as V; V.run(...)

Salida: salidas/validacion/<fecha>/reporte_validacion.xlsx  (el "objeto reporte")

NO escribe nada en la matriz. NO toca los archivos fuente. Solo lee.
"""
from __future__ import annotations

import math
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import yaml

# La consola de Windows suele ser cp1252; evita que un '≤' en un mensaje rompa la corrida.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

# --------------------------------------------------------------------------- rutas
RAIZ = Path(__file__).resolve().parents[1]
CFG_FUENTES = RAIZ / "config" / "fuentes.yaml"
DIR_INSUMOS_DEFECTO = RAIZ / "docs" / "Insumos"
MATRIZ_OBJETIVO_DEFECTO = RAIZ / "docs" / "Objetivo" / "7.SEGUIMIENTO CONTRACTUAL SAVIA PPAL_JULIO.xlsm"
DIR_SALIDAS = RAIZ / "salidas" / "validacion"

MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
            "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# ===========================================================================
#  Normalización
# ===========================================================================
_NA = {None, "", "NA", "N/A", "N.A", "N.A.", "#N/A", "#DIV/0!", "#VALUE!", "SD", "SR", "CR",
       "nan", "NaN", "None", "-", "--"}


_NA_UP = {s.upper() for s in _NA if isinstance(s, str)}


def es_na(x) -> bool:
    if x is None:
        return True
    if isinstance(x, float) and math.isnan(x):
        return True
    if x is pd.NaT or (hasattr(x, "__class__") and x.__class__.__name__ == "NaTType"):
        return True
    if isinstance(x, str):
        return x.strip().upper() in _NA_UP
    return False


def quitar_tildes(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm_texto(x) -> str:
    """MAYÚSCULAS, sin tildes, sin dobles espacios, sin puntuación final."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    s = quitar_tildes(str(x)).upper().strip()
    s = re.sub(r"\s+", " ", s)
    s = s.strip(" .,;:")
    return s


def norm_nit(x) -> str:
    """Deja solo dígitos y quita dígito de verificación tipo '900399132-1'."""
    if x is None:
        return ""
    s = re.sub(r"[^0-9-]", "", str(x))
    s = s.split("-")[0]
    return s.lstrip("0") or s


_RE_CONTRATO = re.compile(r"(?:R[SC]-)?0*(\d{3,5})-?(\d{4})?")


def norm_contrato(x) -> str:
    """
    Normaliza un número de contrato a 'NNNN-AAAA' (sin prefijo de régimen, sin ceros
    de más en la parte numérica pero conservando 4 dígitos). Devuelve '' si no matchea.
    'RS-0603-2026' -> '0603-2026' ; '04449-2025' -> '4449-2025' ; '0351-2025' -> '0351-2025'
    Si la celda trae varios contratos concatenados, toma el primero.
    """
    if x is None:
        return ""
    s = str(x).strip()
    m = _RE_CONTRATO.search(s)
    if not m:
        return ""
    num, anio = m.group(1), m.group(2)
    num = num.zfill(4) if len(num) < 4 else num
    return f"{num}-{anio}" if anio else num


def contratos_en_celda(x) -> list[str]:
    """Todos los contratos que aparezcan en una celda (col D de concurrencia, etc.)."""
    if x is None:
        return []
    out = []
    for m in _RE_CONTRATO.finditer(str(x)):
        num, anio = m.group(1), m.group(2)
        if not anio:
            continue
        num = num.zfill(4) if len(num) < 4 else num
        out.append(f"{num}-{anio}")
    return list(dict.fromkeys(out))


# ===========================================================================
#  Parsing de metas / umbrales / cumplimiento
# ===========================================================================
def parse_umbral(meta):
    """'≤8%'->0.08 ; '>=90%'->0.9 ; '<= 15'->15 ; '2 DÍAS'->2 ; '48 HORAS'->48 ;
       '>95 %'->0.95 ; 'SD'/'NA'->None"""
    if meta is None or es_na(meta):
        return None
    if isinstance(meta, (int, float)):
        return float(meta)
    s = str(meta).strip()
    m = re.search(r"(-?\d+(?:[.,]\d+)?)", s)
    if not m:
        return None
    val = float(m.group(1).replace(",", "."))
    if "%" in s:
        return val / 100 if val > 1 else val
    return val


def operador_de_meta(meta) -> str:
    s = str(meta)
    if "≤" in s or "<=" in s or "<" in s:
        return "<="
    if "≥" in s or ">=" in s or ">" in s:
        return ">="
    return ">="


def a_float(x):
    if isinstance(x, (int, float)) and not (isinstance(x, float) and math.isnan(x)):
        return float(x)
    if x is None:
        return None
    s = str(x).strip().replace("%", "").replace(",", ".")
    try:
        v = float(s)
        return v
    except ValueError:
        return None


def evaluar_cumplimiento(resultado, meta, operador=None, *, ips_en_informe=True,
                         aplica=True, num=None, den=None, componente=None):
    """
    Devuelve (valor_celda, cumplimiento_celda, regla_id) según las 7 variantes de
    config/reglas_cumplimiento.yaml.
    """
    if not aplica:
        return "NA", "NA", "no_aplica_componente"

    meta_s = "" if meta is None else str(meta).strip().upper()
    if meta_s == "SD":
        return (resultado if not es_na(resultado) else "NA"), "Cumple", "meta_SD_cumple_automatico"

    if componente == "oxigeno" and es_na(num) and es_na(den):
        return "NA", "NA", "oxigeno_na_num_y_den"

    if ips_en_informe and es_na(resultado):
        return "NA", "Cumple", "ips_midio_NA_sin_casos"

    if aplica and not ips_en_informe:
        return "NA", "NA", "ips_ausente_del_informe"

    r = a_float(resultado)
    u = parse_umbral(meta)
    op = operador or operador_de_meta(meta)
    if r is None or u is None:
        return (resultado if not es_na(resultado) else "NA"), "NA", "sin_umbral_evaluable"

    if op == "<=":
        ok = r <= u
    elif op == ">=":
        ok = r >= u
    elif op == "==":
        ok = abs(r - u) < 1e-9
    else:
        ok = r <= u
    return resultado, ("Cumple" if ok else "No cumple"), "regla_general"


# ===========================================================================
#  Carga de configuración y de hojas
# ===========================================================================
def cargar_cfg() -> dict:
    with CFG_FUENTES.open(encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def resolver_ph(s, mes_nombre: str, anio: str = "") -> str:
    """Sustituye placeholders de nombre de hoja/archivo por el mes en curso."""
    if not isinstance(s, str):
        return s
    return (s.replace("{MES_NOMBRE_CAP}", mes_nombre.capitalize())
             .replace("{MES_NOMBRE_UP}", mes_nombre.upper())
             .replace("{MES_NOMBRE}", mes_nombre.lower())
             .replace("{ANIO}", str(anio)))


_IGNORAR = ("~$", "7.SEGUIMIENTO CONTRACTUAL SAVIA PPAL")  # temporales de Excel y la propia matriz


def resolver_archivo(carpeta: Path, patron: str) -> Path | None:
    """Busca el patrón en `carpeta` y en TODAS sus subcarpetas (el disco compartido
    tiene los insumos repartidos en subcarpetas: FINANCIERA, MIPRES, AUTORIZACION, ...).
    Si hay varios (p. ej. meses distintos), devuelve el más reciente por fecha de modificación.
    """
    if not patron or not carpeta.exists():
        return None
    hits = [h for h in carpeta.rglob(patron)
            if h.is_file() and not any(h.name.startswith(x) for x in _IGNORAR)]
    hits.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return hits[0] if hits else None


def leer_hoja(ruta: Path, hoja, fila_encabezado: int) -> pd.DataFrame:
    """fila_encabezado es 1-based (como se ve en Excel).

    `keep_default_na=False`: los 'N/A' / 'NA' que el analista escribe LITERALES en los
    insumos se conservan como texto (los interpreta `es_na()` aguas abajo), no como NaN.
    """
    df = pd.read_excel(ruta, sheet_name=hoja, header=fila_encabezado - 1, engine="openpyxl",
                       dtype=object, keep_default_na=False, na_values=[])
    df.columns = [str(c).strip() if c is not None else f"_col{i}" for i, c in enumerate(df.columns)]
    return df


# ===========================================================================
#  Validación de esquema de las fuentes (detección de "drift" a futuro)
# ===========================================================================
@dataclass
class ResultadoFuente:
    componente: str
    estado: str            # OK | FALTA_ARCHIVO | FALTA_HOJA | DRIFT | FALTA_INSUMO | NO_APLICA | MANUAL | SIN_VALIDAR
    archivo: str = ""
    hojas_encontradas: list = field(default_factory=list)
    columnas_faltantes: list = field(default_factory=list)
    columnas_nuevas: list = field(default_factory=list)
    filas: int = 0
    detalle: str = ""


def validar_fuentes(cfg: dict, carpeta: Path, mes_nombre: str = "julio", anio: str = "2026") -> list[ResultadoFuente]:
    from openpyxl import load_workbook

    out: list[ResultadoFuente] = []
    for comp, c in cfg["componentes"].items():
        estado_cfg = (c.get("estado") or "").lower()
        if estado_cfg in ("falta_insumo", "no_aplica", "manual"):
            out.append(ResultadoFuente(comp, estado_cfg.upper(),
                                       detalle=(c.get("nota") or [""])[0] if isinstance(c.get("nota"), list) else ""))
            continue

        # componentes con sub-archivos (domiciliaria, planificación familiar)
        sub = c.get("archivos")
        objetivos = []
        if sub:
            for k, v in sub.items():
                objetivos.append((f"{comp}:{k}", v))
        else:
            objetivos.append((comp, c))

        for nombre, spec in objetivos:
            patron = spec.get("archivo_glob")
            ruta = resolver_archivo(carpeta, patron) if patron else None
            if patron and ruta is None:
                out.append(ResultadoFuente(nombre, "FALTA_ARCHIVO", detalle=f"glob={patron!r}"))
                continue
            if ruta is None:
                out.append(ResultadoFuente(nombre, "SIN_VALIDAR", detalle="sin archivo_glob"))
                continue

            hojas_pedidas = spec.get("hojas") or ([spec["hoja"]] if spec.get("hoja") else None)
            if hojas_pedidas:
                hojas_pedidas = [resolver_ph(h, mes_nombre, anio) for h in hojas_pedidas]
            try:
                wb = load_workbook(ruta, read_only=True, data_only=True, keep_vba=False)
                hojas_libro = wb.sheetnames
                wb.close()
            except Exception as e:  # noqa: BLE001
                out.append(ResultadoFuente(nombre, "FALTA_ARCHIVO", archivo=ruta.name,
                                           detalle=f"no se pudo abrir: {e}"))
                continue

            if hojas_pedidas and hojas_pedidas != ["*"]:
                faltan_hojas = [h for h in hojas_pedidas if h not in hojas_libro and h != "*"]
                if faltan_hojas:
                    out.append(ResultadoFuente(nombre, "FALTA_HOJA", archivo=ruta.name,
                                               hojas_encontradas=hojas_libro,
                                               detalle=f"faltan hojas: {faltan_hojas}"))
                    continue

            esperadas = spec.get("columnas_esperadas") or []
            if not esperadas:
                out.append(ResultadoFuente(nombre, "OK", archivo=ruta.name,
                                           hojas_encontradas=hojas_pedidas or hojas_libro,
                                           detalle="sin lista de columnas esperadas -> solo se validó archivo/hoja"))
                continue

            hoja0 = (hojas_pedidas or hojas_libro)[0]
            if hoja0 == "*":
                hoja0 = hojas_libro[0]
            fe = spec.get("fila_encabezado") or 1
            try:
                df = leer_hoja(ruta, hoja0, fe)
            except Exception as e:  # noqa: BLE001
                out.append(ResultadoFuente(nombre, "DRIFT", archivo=ruta.name,
                                           detalle=f"no se pudo leer hoja {hoja0!r} fila {fe}: {e}"))
                continue

            cols_norm = [norm_texto(x) for x in df.columns]
            faltan = [e for e in esperadas
                      if not any(norm_texto(e) == cn or cn.startswith(norm_texto(e)) for cn in cols_norm)]
            estado = "DRIFT" if faltan else "OK"
            out.append(ResultadoFuente(
                nombre, estado, archivo=ruta.name,
                hojas_encontradas=hojas_pedidas or [hoja0],
                columnas_faltantes=faltan,
                filas=int(len(df)),
                detalle="" if not faltan else f"{len(faltan)} columna(s) esperada(s) no encontrada(s)",
            ))
    return out


# ===========================================================================
#  Lectura de la matriz objetivo (General)  ->  wide por contrato
# ===========================================================================
from openpyxl.utils import get_column_letter, column_index_from_string  # noqa: E402


def leer_general(ruta_xlsm: Path) -> pd.DataFrame:
    """
    Devuelve un DataFrame indexado por número de contrato normalizado, con columnas
    'A'..'NO' (letras de Excel) y el valor de cada celda de la hoja General.
    Fila 2 = encabezados; fila 3+ = datos.
    """
    from openpyxl import load_workbook

    wb = load_workbook(ruta_xlsm, read_only=True, data_only=True, keep_vba=False)
    ws = wb["General"]
    filas = []
    idx = []
    for r, fila in enumerate(ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True), start=3):
        if fila[0] in (None, ""):
            continue
        d = {}
        for c, v in enumerate(fila, start=1):
            d[get_column_letter(c)] = v
        filas.append(d)
        idx.append(norm_contrato(fila[0]))
    wb.close()
    df = pd.DataFrame(filas, index=idx)
    df.index.name = "numero_contrato"
    return df


def encabezados_general(ruta_xlsm: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    wb = load_workbook(ruta_xlsm, read_only=True, data_only=True, keep_vba=False)
    ws = wb["General"]
    fila2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    wb.close()
    return {get_column_letter(i): (str(v).strip() if v is not None else "")
            for i, v in enumerate(fila2, start=1)}


def indice_llaves(general: pd.DataFrame) -> dict:
    """Mapas de cruce derivados de la propia matriz objetivo (col B=NIT, F=REGIMEN, K=MUNICIPIO)."""
    nit_a_contratos: dict[str, list[str]] = {}
    muni_a_contratos: dict[str, list[str]] = {}
    for contrato, row in general.iterrows():
        nit = norm_nit(row.get("B"))
        if nit:
            nit_a_contratos.setdefault(nit, []).append(contrato)
        muni = norm_texto(row.get("K"))
        if muni:
            muni_a_contratos.setdefault(muni, []).append(contrato)
    return {
        "nit_a_contratos": nit_a_contratos,
        "muni_a_contratos": muni_a_contratos,
        "regimen": {c: norm_texto(r.get("F")) for c, r in general.iterrows()},
        "modalidad": {c: norm_texto(r.get("H")) for c, r in general.iterrows()},
        "categoria": {c: norm_texto(r.get("I")) for c, r in general.iterrows()},
        "contratos": set(general.index),
    }


# ===========================================================================
#  Extractores por componente  ->  filas {numero_contrato, col, valor}
# ===========================================================================
def _rows(*triples):
    return [{"numero_contrato": a, "col": b, "valor": c} for a, b, c in triples if a]


# --- Datos generales: SIRECI -> columnas A:AE de General (migración directa) ---
MAPEO_SIRECI_GENERAL = {
    "A": "NUMERO DE CONTRATO",
    "B": "CONTRATISTA : NUMERO DEL NIT",
    "C": "REPS",
    "D": "CONTRATISTA : NOMBRE COMPLETO",
    "E": "OBJETO DEL CONTRATO",
    "F": "REGIMEN",
    "G": "NATURALEZA JURIDICA: RED PUBLICA / PRIVADA",
    "H": "MODALIDAD",
    "I": "CATEGORIA DEL CONTRATO SEGUIMIENTO A LA RED",
    "J": "SUBREGION",
    "K": "MUNICIPIO",
    "L": "VALOR INICIAL DEL CONTRATO En pesos",
    "M": "ADICIONES : VALOR TOTAL",
    "N": "VALOR TOTAL INCLUIDA ADICION  En pesos",
    # O/P/Q (PLAZO INICIAL / TIEMPO PRÓRROGAS / INCLUIDA PRÓRROGA): el objetivo los deja
    # casi siempre en blanco -> no se migran para no ensuciar la comparación.
    "R": "FECHA INICIO CONTRATO",
    "S": "FECHA TERMINACION CONTRATO",
    "T": "TOTAL FACTURACION ACUMULADA",
    "U": "PORCENTAJE AVANCE PRESUPUESTAL PROGRAMADO",
    "V": "PORCENTAJE AVANCE PRESUPUESTAL REAL",
    "W": "%  SOBREEJECUCION / SUBEJECUCION",
    "X": "ESTADO DE LEGALIZACION",
    "Y": "CORTE SIRECI",
    "Z": "POLIZA CUMPLIMIENTO VIGENCIA DESDE",
    "AA": "POLIZA CUMPLIMIENTO  VIGENCIA HASTA",
    "AB": "POLIZA CUMPLIMIENTO",
    "AC": "POLIZA RC  VIGENCIA DESDE",
    "AD": "POLIZA RC  VIGENCIA HASTA",
    "AE": "POLIZA RC",
}


def ex_datos_generales(ruta, cfg, idx=None) -> list[dict]:
    """Migración directa SIRECI -> A:AE. Además define el ORDEN de filas de General
    (el objetivo de julio está en el mismo orden que SIRECI)."""
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    cols = {norm_texto(c): c for c in df.columns}
    out = []
    for _, r in df.iterrows():
        contrato = norm_contrato(r.get(cols.get(norm_texto("NUMERO DE CONTRATO"), "NUMERO DE CONTRATO")))
        if not contrato:
            continue
        for gcol, scol in MAPEO_SIRECI_GENERAL.items():
            real = cols.get(norm_texto(scol))
            if real is None:
                continue
            out.append({"numero_contrato": contrato, "col": gcol, "valor": r.get(real)})
    return out


def orden_contratos_sireci(ruta, cfg) -> list[str]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    cols = {norm_texto(c): c for c in df.columns}
    col = cols.get(norm_texto("NUMERO DE CONTRATO"), "NUMERO DE CONTRATO")
    vistos, orden = set(), []
    for _, r in df.iterrows():
        c = norm_contrato(r.get(col))
        if c and c not in vistos:
            vistos.add(c)
            orden.append(c)
    return orden


def ex_auditoria_calidad(ruta, cfg, idx) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    out = []
    for _, r in df.iterrows():
        nit = norm_nit(r.get("NIT"))
        for contrato in idx["nit_a_contratos"].get(nit, []):
            out += _rows((contrato, "AF", r.get("FECHA")), (contrato, "AG", r.get("RESULTADO")))
    return out


def ex_financieros_evento(ruta, cfg, idx, mes_nombre) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    per = df["Periodo"].astype(str).str.strip().str.upper() == mes_nombre.upper()
    df = df[per]
    out = []
    m = {"AH": "Meta glosas", "AI": "Indicador glosa", "AJ": "Cumplimiento glosa",
         "AK": "Meta devoluciones", "AL": "Indicador devoluciones", "AM": "Cumplimiento Devoluciones",
         "AN": "Meta", "AO": "Indicador radicación ≤ 180 días", "AP": "Cumplimiento"}
    for _, r in df.iterrows():
        nit = norm_nit(r.get("NIT_IPS"))
        contratos = idx["nit_a_contratos"].get(nit, [])
        reg_fila = norm_texto(r.get("Regimen"))
        for contrato in contratos:
            reg_c = idx["regimen"].get(contrato, "")
            # si el contrato es de un solo régimen y la fila es de otro, saltar
            if reg_c in ("SUBSIDIADO", "CONTRIBUTIVO") and reg_fila and reg_fila != reg_c:
                continue
            if idx["modalidad"].get(contrato) not in ("EVENTO", ""):
                continue
            for col, campo in m.items():
                out.append({"numero_contrato": contrato, "col": col, "valor": r.get(campo)})
    return out


def ex_concurrencia(ruta, cfg, idx) -> list[dict]:
    out = []
    for hoja in cfg["hojas"]:
        df = leer_hoja(ruta, hoja, cfg["fila_encabezado"])
        cols = {norm_texto(c): c for c in df.columns}
        c_meta_re = cols.get(norm_texto("META  PROPORCION REINGRESOS ANTES DE 15 DIAS"))
        c_ind_re = cols.get(norm_texto("INDICADOR  PROPORCION REINGRESOS ANTES DE 15 DIAS"))
        c_meta_de = cols.get(norm_texto("META   PROMEDIO DIAS ESTANCIA POR EVENTO"))
        c_ind_de = cols.get(norm_texto("INDICADOR  PROMEDIO DIAS ESTANCIA POR EVENTO"))
        for _, r in df.iterrows():
            nit = norm_nit(r.get("NIT"))
            for contrato in idx["nit_a_contratos"].get(nit, []):
                out += _rows(
                    (contrato, "AZ", r.get(c_meta_re)), (contrato, "BA", r.get(c_ind_re)),
                    (contrato, "BC", r.get(c_meta_de)), (contrato, "BD", r.get(c_ind_de)),
                )
    return out


def ex_mipres(ruta_ent, ruta_jun, cfg_ent, cfg_jun, idx) -> list[dict]:
    out = []
    if ruta_ent is not None:
        df = leer_hoja(ruta_ent, cfg_ent["hoja"], cfg_ent["fila_encabezado"])
        # header en fila 5: C/D = CONTRIBUTIVO num/den ; E/F = SUBSIDIADO num/den ; G/H = (en blanco)
        num_cols = [c for c in df.columns if c.lower().startswith("numerador")]
        den_cols = [c for c in df.columns if c.lower().startswith("denominador")]
        reg_par = {"CONTRIBUTIVO": (num_cols[0:1], den_cols[0:1]),
                   "SUBSIDIADO": (num_cols[1:2], den_cols[1:2])}
        agg: dict[str, dict[str, float]] = {}
        for _, r in df.iterrows():
            nit = norm_nit(r.get("Nit_del_proveedor"))
            a = agg.setdefault(nit, {"CONTRIBUTIVO_n": 0.0, "CONTRIBUTIVO_d": 0.0,
                                     "SUBSIDIADO_n": 0.0, "SUBSIDIADO_d": 0.0})
            for reg, (ncs, dcs) in reg_par.items():
                a[f"{reg}_n"] += sum(a_float(r.get(c)) or 0 for c in ncs)
                a[f"{reg}_d"] += sum(a_float(r.get(c)) or 0 for c in dcs)
        for nit, a in agg.items():
            for contrato in idx["nit_a_contratos"].get(nit, []):
                reg = idx["regimen"].get(contrato, "")
                if reg == "SUBSIDIADO":
                    n, d = a["SUBSIDIADO_n"], a["SUBSIDIADO_d"]
                elif reg == "CONTRIBUTIVO":
                    n, d = a["CONTRIBUTIVO_n"], a["CONTRIBUTIVO_d"]
                else:
                    n = a["SUBSIDIADO_n"] + a["CONTRIBUTIVO_n"]
                    d = a["SUBSIDIADO_d"] + a["CONTRIBUTIVO_d"]
                out.append({"numero_contrato": contrato, "col": "BG", "valor": (n / d) if d else None})
    if ruta_jun is not None:
        df = leer_hoja(ruta_jun, cfg_jun["hoja"], cfg_jun["fila_encabezado"])
        for _, r in df.iterrows():
            nit = norm_nit(r.get("n_identificacion_ips"))
            n, d = a_float(r.get("Numerador")), a_float(r.get("Denominador"))
            val = (n / d) if d else None
            for contrato in idx["nit_a_contratos"].get(nit, []):
                out.append({"numero_contrato": contrato, "col": "BJ", "valor": val})
    return out


def _autoriz(ruta, cfg, idx, cols_por_regimen, col_general, mes_nombre):
    df = leer_hoja(ruta, resolver_ph(cfg["hoja"], mes_nombre), cfg["fila_encabezado"])
    # tras leer con header=fila 2, las 3 apariciones quedan como  'X', 'X.1', 'X.2'
    base = cols_por_regimen
    variantes = [base, f"{base}.1", f"{base}.2"]  # SUBSIDIADO / CONTRIBUTIVO / TOTAL
    out = []
    nit_col = [c for c in df.columns if c.upper().startswith("NIT_PRESTADOR")][0]
    for _, r in df.iterrows():
        nit = norm_nit(r.get(nit_col))
        for contrato in idx["nit_a_contratos"].get(nit, []):
            reg = idx["regimen"].get(contrato, "")
            v = variantes[0] if reg == "SUBSIDIADO" else variantes[1] if reg == "CONTRIBUTIVO" else variantes[2]
            if v not in df.columns:
                v = variantes[2] if variantes[2] in df.columns else variantes[0]
            out.append({"numero_contrato": contrato, "col": col_general, "valor": r.get(v)})
    return out


def ex_transporte(ruta, cfg, idx) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    cols = {norm_texto(c): c for c in df.columns}
    c_meta = [c for c in df.columns if norm_texto(c) == "META"]
    c_ind1 = cols.get(norm_texto("INDICADOR Proporción de traslados asistenciales efectivos"))
    c_ind2 = cols.get(norm_texto("INDICADOR Cumplimiento Tiempo Promesa de Servicio (Minutos)"))
    out = []
    for _, r in df.iterrows():
        contrato = norm_contrato(r.get("NÚMERO DE CONTRATO") or r.get("NUMERO DE CONTRATO"))
        if contrato not in idx["contratos"]:
            continue
        meta1 = r.get(c_meta[0]) if c_meta else None
        meta2 = r.get(c_meta[1]) if len(c_meta) > 1 else None
        out += _rows((contrato, "DQ", meta1), (contrato, "DR", r.get(c_ind1)),
                     (contrato, "DT", meta2), (contrato, "DU", r.get(c_ind2)))
    return out


def ex_tutelas(ruta, cfg, idx) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    out = []
    for _, r in df.iterrows():
        contrato = norm_contrato(r.get("Contrato"))
        if contrato not in idx["contratos"]:
            continue
        out.append({"numero_contrato": contrato, "col": "JC", "valor": r.get("Indicador")})
    return out


def ex_extramuralidad(ruta, cfg, idx) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    out = []
    for _, r in df.iterrows():
        muni = norm_texto(r.get("MUNICIPIO"))
        for contrato in idx["muni_a_contratos"].get(muni, []):
            if idx["modalidad"].get(contrato) not in ("CAPITA", "CÁPITA", ""):
                continue
            out += _rows((contrato, "NI", r.get("META INDICADOR")),
                         (contrato, "NJ", r.get("RESULTADO DEL INDICADOR")))
    return out


def ex_medicamentos(ruta, cfg, idx) -> list[dict]:
    df = leer_hoja(ruta, cfg["hoja"], cfg["fila_encabezado"])
    cols = {norm_texto(c): c for c in df.columns}

    def g(nombre):
        return cols.get(norm_texto(nombre))

    m = {
        "FM": g("A14 - META: Promedio de tiempo de espera para la entrega de medicamentos incluidos en el PBS (días)"),
        "FN": g("A14 - Resultado: Promedio de tiempo de espera para la entrega de medicamentos incluidos en el PBS"),
        "FP": g("A15 - META: Porcentaje  de fórmulas médicas entregadas de manera completa"),
        "FQ": g("A15 - Resultado: Porcentaje  de fórmulas médicas entregadas de manera completa"),
        "FS": g("A16 - META: Porcentaje de fórmulas médicas entregadas de manera oportuna"),
        "FT": g("A16 - Resultado: Porcentaje de fórmulas médicas entregadas de manera oportuna"),
    }
    out = []
    for _, r in df.iterrows():
        contrato = norm_contrato(r.get("NÚMERO DE CONTRATO") or r.get("NUMERO DE CONTRATO"))
        if contrato not in idx["contratos"]:
            continue
        for col, campo in m.items():
            if campo:
                out.append({"numero_contrato": contrato, "col": col, "valor": r.get(campo)})
    return out


# ===========================================================================
#  Comparación con la matriz objetivo
# ===========================================================================
def _norm_cmp(v):
    """Normaliza un valor para comparar proceso vs objetivo."""
    if es_na(v):
        return "NA"
    f = a_float(v)
    if f is not None:
        return round(f, 4)
    return norm_texto(v)


def comparar(extraido: list[dict], general: pd.DataFrame, enc: dict) -> pd.DataFrame:
    filas = []
    for e in extraido:
        c, col, vp = e["numero_contrato"], e["col"], e["valor"]
        vo = general.at[c, col] if (c in general.index and col in general.columns) else None
        a, b = _norm_cmp(vp), _norm_cmp(vo)
        if isinstance(a, float) and isinstance(b, float):
            igual = abs(a - b) <= max(1e-4, abs(b) * 1e-3)
        else:
            igual = a == b
        clasif = "match"
        if not igual:
            if b == "NA" and a != "NA":
                clasif = "objetivo_NA_proceso_valor"
            elif a == "NA" and b != "NA":
                clasif = "llave_no_cruzada_o_dato_faltante"
            elif isinstance(a, float) and isinstance(b, float):
                clasif = "diferencia_numerica"
            else:
                clasif = "diferencia_texto"
        filas.append({
            "numero_contrato": c, "col": col, "encabezado": enc.get(col, "")[:60],
            "valor_proceso": vp, "valor_objetivo": vo,
            "cmp_proceso": a, "cmp_objetivo": b, "igual": igual, "clasificacion": clasif,
        })
    return pd.DataFrame(filas)


# ===========================================================================
#  Orquestación
# ===========================================================================
BANDAS = {
    "DATOS GENERALES": ("A", "AE"), "AUDITORÍA CALIDAD": ("AF", "AG"), "FINANCIEROS": ("AH", "AP"),
    "INDICADORES DE ACCESO": ("AQ", "AY"), "CONCURRENCIA": ("AZ", "BE"), "MIPRES": ("BF", "BK"),
    "MEDICINA DOMICILIARIA": ("BL", "CX"), "SALUD ORAL": ("CY", "DP"), "TRANSPORTE ASISTENCIAL": ("DQ", "DV"),
    "HOGARES DE PASO": ("DW", "EE"), "OXÍGENO": ("EF", "FC"), "MOS": ("FD", "FL"),
    "MEDICAMENTOS FÉNIX": ("FM", "GA"), "REPORTE 1552": ("GB", "JA"), "TUTELAS": ("JB", "JD"),
    "PQRSD": ("JE", "JV"), "SALUD MENTAL": ("JW", "KB"), "PLANIFICACIÓN FAMILIAR": ("KC", "LC"),
    "AYUDAS DIAGNÓSTICAS": ("LD", "NH"), "Campañas extramural": ("NI", "NK"),
}


def banda_de_col(col: str) -> str:
    ci = column_index_from_string(col)
    for nombre, (a, b) in BANDAS.items():
        if column_index_from_string(a) <= ci <= column_index_from_string(b):
            return nombre
    return "?"


def run(mes: str = "2026-07",
        ruta_unidad: Path | str | None = None,
        carpeta_insumos: Path | str | None = None,
        subcarpeta_insumos: str = "",
        matriz_objetivo: Path | str | None = MATRIZ_OBJETIVO_DEFECTO,
        salida: Path | str | None = None) -> dict:
    """
    Parámetros
    ----------
    ruta_unidad : raíz del disco compartido (p. ej. r"Z:\\10.INDICADORES SEGUIMIENTO CONTRACTUAL").
                  Si se da y no se da `carpeta_insumos`, los insumos se buscan aquí
                  (recursivamente, en todas las subcarpetas).
    carpeta_insumos : carpeta concreta con los insumos. Tiene prioridad sobre `ruta_unidad`.
    subcarpeta_insumos : subcarpeta bajo `ruta_unidad` (opcional, p. ej. el mes).
    matriz_objetivo : .xlsm diligenciado para la comparación. Si no existe, se omite la
                      comparación y solo se valida el esquema de las fuentes.

    Nunca lanza excepción por rutas inexistentes: lo reporta y sigue con lo que pueda.
    """
    anio, m = mes.split("-")
    mes_nombre = MESES_ES[int(m) - 1]
    cfg = cargar_cfg()

    # ---- resolver la carpeta de insumos -------------------------------------
    if carpeta_insumos:
        carpeta = Path(carpeta_insumos)
    elif ruta_unidad:
        carpeta = Path(ruta_unidad) / subcarpeta_insumos if subcarpeta_insumos else Path(ruta_unidad)
    else:
        carpeta = DIR_INSUMOS_DEFECTO

    fecha = datetime.now().strftime("%Y-%m-%d_%H%M")
    out_dir = Path(salida) if salida else (DIR_SALIDAS / fecha)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / "reporte_validacion.xlsx"

    if not carpeta.exists():
        msg = (f"NO se encontró la carpeta de insumos: {carpeta}\n"
               f"       Revisa que la unidad compartida esté conectada y que la ruta "
               f"en RUTA_UNIDAD / CARPETA_INSUMOS sea correcta.")
        print("[ERROR] " + msg)
        err = pd.DataFrame([["carpeta_insumos_no_encontrada", str(carpeta)]],
                           columns=["problema", "detalle"])
        with pd.ExcelWriter(xlsx, engine="xlsxwriter") as xw:
            err.to_excel(xw, sheet_name="ERROR", index=False)
        return {"reporte": xlsx, "error": msg, "fuentes": pd.DataFrame(),
                "concordancia_detalle": pd.DataFrame(), "concordancia_por_banda": pd.DataFrame(),
                "cobertura": pd.DataFrame(), "resumen": err, "faltantes": pd.DataFrame(),
                "que_revisar": pd.DataFrame()}

    print(f"[1/5] Validando esquema de fuentes en {carpeta} ...")
    fuentes = validar_fuentes(cfg, carpeta, mes_nombre, anio)
    df_fuentes = pd.DataFrame([r.__dict__ for r in fuentes])

    # ---- lista de fuentes NO encontradas / con problema --------------------
    _falla = {"FALTA_ARCHIVO", "FALTA_HOJA", "FALTA_INSUMO", "DRIFT"}
    faltantes = []
    for _, fr in df_fuentes.iterrows():
        if fr["estado"] not in _falla:
            continue
        spec = cfg["componentes"].get(fr["componente"].split(":")[0], {})
        faltantes.append({
            "componente": fr["componente"],
            "estado": fr["estado"],
            "archivo_esperado": spec.get("archivo_glob") or spec.get("fuente") or "(sub-archivo)",
            "hoja_esperada": spec.get("hoja") or ", ".join(spec.get("hojas", []) or []),
            "archivo_encontrado": fr["archivo"] or "-",
            "columnas_faltantes": ", ".join(fr["columnas_faltantes"] or []),
            "detalle": fr["detalle"],
            "carpeta_buscada": str(carpeta),
        })
    df_faltantes = pd.DataFrame(faltantes)
    if not df_faltantes.empty:
        print(f"       !! {len(df_faltantes)} fuente(s) con problema "
              f"(ver hoja 'faltantes' del reporte):")
        for f in faltantes:
            print(f"         - {f['componente']:26s} {f['estado']:13s} {f['detalle']}")
    else:
        print("       Todas las fuentes esperadas se encontraron y su esquema coincide.")

    # ---- matriz objetivo (opcional) --------------------------------------
    obj = Path(matriz_objetivo) if matriz_objetivo else None
    hay_objetivo = bool(obj and obj.exists())
    if not hay_objetivo:
        print(f"[2/5] Matriz objetivo no disponible ({obj}) -> se omite la comparación.")
        general = pd.DataFrame()
        enc, idx = {}, {"nit_a_contratos": {}, "muni_a_contratos": {}, "regimen": {},
                        "modalidad": {}, "categoria": {}, "contratos": set()}
    else:
        print(f"[2/5] Leyendo matriz objetivo {obj.name} ...")
        general = leer_general(obj)
        enc = encabezados_general(obj)
        idx = indice_llaves(general)
        print(f"       {len(general)} contratos, {general.notna().sum().sum():,} celdas con dato.")

    print("[3/5] Extrayendo componentes reconstruibles ...")
    C = cfg["componentes"]
    extr: list[dict] = []
    log = []

    def intentar(nombre, fn):
        try:
            filas = fn()
            extr.extend(filas)
            log.append((nombre, "ok", len(filas)))
            print(f"       {nombre:24s} -> {len(filas)} celdas")
        except Exception as e:  # noqa: BLE001
            log.append((nombre, f"error: {e}", 0))
            print(f"       {nombre:24s} -> ERROR {e}")

    def ruta(comp):
        return resolver_archivo(carpeta, C[comp].get("archivo_glob", "")) if comp in C else None

    if (r := ruta("auditoria_calidad")):
        intentar("auditoria_calidad", lambda: ex_auditoria_calidad(r, C["auditoria_calidad"], idx))
    if (r := ruta("financieros_evento")):
        intentar("financieros_evento", lambda: ex_financieros_evento(r, C["financieros_evento"], idx, mes_nombre))
    if (r := ruta("concurrencia")):
        intentar("concurrencia", lambda: ex_concurrencia(r, C["concurrencia"], idx))
    re_, rj_ = ruta("mipres_entregas"), ruta("mipres_juntas")
    if re_ or rj_:
        intentar("mipres", lambda: ex_mipres(re_, rj_, C["mipres_entregas"], C["mipres_juntas"], idx))
    if (r := ruta("autorizaciones_tramite")):
        intentar("autorizaciones_tramite",
                 lambda: _autoriz(r, C["autorizaciones_tramite"], idx, "% TRÁMITE", "AU", mes_nombre))
    if (r := ruta("autorizaciones_marcacion")):
        intentar("autorizaciones_marcacion",
                 lambda: _autoriz(r, C["autorizaciones_marcacion"], idx, "% FECHA PRESTACIÓN", "AX", mes_nombre))
    if (r := ruta("transporte_asistencial")):
        intentar("transporte_asistencial", lambda: ex_transporte(r, C["transporte_asistencial"], idx))
    if (r := ruta("tutelas")):
        intentar("tutelas", lambda: ex_tutelas(r, C["tutelas"], idx))
    if (r := ruta("extramuralidad")):
        intentar("extramuralidad", lambda: ex_extramuralidad(r, C["extramuralidad"], idx))
    if (r := ruta("medicamentos")):
        intentar("medicamentos", lambda: ex_medicamentos(r, C["medicamentos"], idx))

    if hay_objetivo:
        print("[4/5] Comparando contra la matriz objetivo ...")
        df_cmp = comparar(extr, general, enc)
    else:
        print("[4/5] (sin matriz objetivo: no hay comparación)")
        df_cmp = pd.DataFrame()
    if not df_cmp.empty:
        df_cmp["banda"] = df_cmp["col"].map(banda_de_col)
        resumen = (df_cmp.groupby("banda")
                   .agg(celdas=("igual", "size"), coinciden=("igual", "sum"))
                   .assign(pct=lambda d: (100 * d.coinciden / d.celdas).round(1))
                   .reset_index())
        clasif = df_cmp.groupby("clasificacion").size().reset_index(name="celdas").sort_values("celdas", ascending=False)
    else:
        resumen = pd.DataFrame()
        clasif = pd.DataFrame()

    # cobertura de la matriz: cuántas celdas de indicadores tiene el objetivo por banda
    cob = []
    if hay_objetivo:
        for banda, (a, b) in BANDAS.items():
            ca, cb = column_index_from_string(a), column_index_from_string(b)
            letras = [get_column_letter(i) for i in range(ca, cb + 1) if get_column_letter(i) in general.columns]
            sub = general[letras]
            con_dato = int((sub.notna() & (sub != "NA")).sum().sum())
            cob.append({"banda": banda, "columnas": len(letras),
                        "celdas_objetivo_con_valor_no_NA": con_dato,
                        "reconstruible_ahora": banda in set(df_cmp["banda"]) if not df_cmp.empty else False})
    df_cob = pd.DataFrame(cob)

    # discrepancias Indicaciones.txt vs realidad (fijas, halladas en la re-evaluación)
    disc = pd.DataFrame([
        ["concurrencia", "hoja", "IPS_Concurrencia_CONEXIONES", "IPS_Concurrencia_CONEXIONBES",
         "El archivo real trae 'CONEXIONBES' (typo). Usar el nombre real."],
        ["transporte_asistencial", "hoja", "Indicaddores_Contratos_Evento", "Indicadores_Contratos_Evento",
         "Indicaciones.txt tiene un typo ('dd'). Hoja real: 'Indicadores_Contratos_Evento'."],
        ["medicina_domiciliaria (curativa)", "hoja", "CONTRATO 0426 2025",
         "CONTRATO 0227 2024 + CONTRATO 0429 2025",
         "El archivo Curativ trae 2 hojas y ninguna es '0426 2025'."],
        ["financieros", "archivo/hoja", "fuentes.yaml: '9. Consolidado indicadores Financieros por IPS', hoja null",
         "9.Consolidado indicadores por IPS <MES> <AÑO>, hoja '2026'",
         "Ajustar fuentes.yaml al nombre/hoja reales."],
        ["nota_tecnica_capita", "insumo", "Tablero de Seguimiento Ejecucion Capitas ENE-JUL 2026 / hoja resumen",
         "(ausente)", "No está en docs/Insumos. Falta para INDICADORES DE ACCESO (col AR)."],
        ["financieros_capita / financieros_pgp / pqrsd", "insumo", "INFORME CAPITA Y MOVILIDAD; 2 Google Sheets",
         "(ausentes)", "Requieren acceso a Drive / carpeta compartida."],
        ["salud_mental", "alcance", "banda JW:KB", "todo NA",
         "Indicaciones.txt: 'Ya no mandan indicadores'."],
        ["auditoria_calidad", "hoja", "(fuentes.yaml pide 'fecha de auditoría', '% de calificación')",
         "hoja 'Indicador' con columnas NIT/CONTRATISTA/FECHA/RESULTADO",
         "Indicaciones.txt manda la hoja simple 'Indicador'."],
    ], columns=["componente", "tipo", "decia", "es_en_realidad", "accion"])

    # ---- "qué revisar": lista de acciones consolidada para el analista -----
    revisar = []
    for f in faltantes:
        prio = "ALTA" if f["estado"] in ("FALTA_ARCHIVO", "FALTA_HOJA") else "MEDIA"
        revisar.append({"prioridad": prio, "tema": f"fuente {f['estado']}",
                        "componente": f["componente"],
                        "que_hacer": f["detalle"] or f"Conseguir/revisar {f['archivo_esperado']}"})
    if not resumen.empty:
        for _, rr in resumen.iterrows():
            if rr["pct"] < 95:
                revisar.append({"prioridad": "MEDIA", "tema": f"concordancia {rr['pct']}%",
                                "componente": rr["banda"],
                                "que_hacer": "Revisar hoja 'concordancia_detalle' filtrando esta banda "
                                             "(diferencias de dato / regla / llave)."})
    df_revisar = pd.DataFrame(revisar) if revisar else pd.DataFrame(
        [{"prioridad": "-", "tema": "sin observaciones", "componente": "-",
          "que_hacer": "Todas las fuentes OK y concordancia >= 95% donde hay extractor."}])

    resumen_global = pd.DataFrame([
        ["mes", mes],
        ["carpeta_insumos", str(carpeta)],
        ["matriz_objetivo", str(obj) if hay_objetivo else "(no disponible)"],
        ["contratos_en_matriz", len(general) if hay_objetivo else "-"],
        ["fuentes_esperadas", len(df_fuentes)],
        ["fuentes_OK", int((df_fuentes.estado == "OK").sum())],
        ["fuentes_DRIFT", int((df_fuentes.estado == "DRIFT").sum())],
        ["fuentes_FALTA_ARCHIVO", int((df_fuentes.estado == "FALTA_ARCHIVO").sum())],
        ["fuentes_FALTA_HOJA", int((df_fuentes.estado == "FALTA_HOJA").sum())],
        ["fuentes_FALTA_INSUMO", int((df_fuentes.estado == "FALTA_INSUMO").sum())],
        ["celdas_reconstruidas", len(df_cmp)],
        ["celdas_coinciden", int(df_cmp["igual"].sum()) if not df_cmp.empty else 0],
        ["pct_concordancia_global",
         round(100 * df_cmp["igual"].mean(), 1) if not df_cmp.empty else None],
    ], columns=["metrica", "valor"])

    print(f"[5/5] Escribiendo objeto reporte -> {xlsx}")
    with pd.ExcelWriter(xlsx, engine="xlsxwriter") as xw:
        resumen_global.to_excel(xw, sheet_name="resumen", index=False)
        df_revisar.to_excel(xw, sheet_name="que_revisar", index=False)
        (df_faltantes if not df_faltantes.empty else
         pd.DataFrame([{"info": "ninguna fuente esperada faltó ni cambió de formato"}])
         ).to_excel(xw, sheet_name="faltantes", index=False)
        df_fuentes.to_excel(xw, sheet_name="fuentes_esquema", index=False)
        disc.to_excel(xw, sheet_name="discrepancias_indicaciones", index=False)
        if not df_cob.empty:
            df_cob.to_excel(xw, sheet_name="cobertura_matriz", index=False)
        if not resumen.empty:
            resumen.to_excel(xw, sheet_name="concordancia_por_banda", index=False)
        if not clasif.empty:
            clasif.to_excel(xw, sheet_name="concordancia_clasificacion", index=False)
        if not df_cmp.empty:
            df_cmp.sort_values(["banda", "numero_contrato", "col"]).to_excel(
                xw, sheet_name="concordancia_detalle", index=False)
        pd.DataFrame(log, columns=["componente", "estado", "celdas"]).to_excel(
            xw, sheet_name="log_extraccion", index=False)

    return {
        "reporte": xlsx, "fuentes": df_fuentes, "concordancia_detalle": df_cmp,
        "concordancia_por_banda": resumen, "cobertura": df_cob, "resumen": resumen_global,
        "faltantes": df_faltantes, "que_revisar": df_revisar,
    }


if __name__ == "__main__":
    res = run()
    print("\n=== RESUMEN ===")
    print(res["resumen"].to_string(index=False))
    print("\n=== FUENTES ===")
    print(res["fuentes"][["componente", "estado", "archivo", "filas", "detalle"]].to_string(index=False))
    if not res["concordancia_por_banda"].empty:
        print("\n=== CONCORDANCIA POR BANDA ===")
        print(res["concordancia_por_banda"].to_string(index=False))
