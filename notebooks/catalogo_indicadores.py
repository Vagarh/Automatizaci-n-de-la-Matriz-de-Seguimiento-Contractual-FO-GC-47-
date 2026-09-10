"""Genera el catálogo columna-por-columna de la hoja `General` leyendo su fila 2.

Uso:
    python notebooks/catalogo_indicadores.py --matriz "ruta\\a\\matriz.xlsm"
    python notebooks/catalogo_indicadores.py --matriz ... --yaml > config/indicadores.generado.yaml

Detecta el patrón de TRIPLETAS META / RESULTADO / CUMPLIMIENTO por banda de grupo
(fila 1) y emite, por cada indicador, sus 3 columnas destino.

Heurística de rol de columna (por el texto del encabezado, fila 2):
    META*        -> meta
    CUMPL* / CUMPLIMIENTO*  -> cumplimiento
    resto        -> resultado
Las bandas 1552 usan 'PFAFS' como columna de resultado.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HOJA = "General"
FILA_GRUPOS = 1
FILA_HEADERS = 2


def _rol(header: str) -> str:
    h = (header or "").strip().upper()
    if h.startswith("META") or h.startswith("≥") or h.startswith("≤") or " META" in h[:12]:
        return "meta"
    if h.startswith("CUMPL"):
        return "cumplimiento"
    return "resultado"


def leer_catalogo(ruta_xlsm: Path) -> list[dict]:
    wb = load_workbook(ruta_xlsm, read_only=True, data_only=True, keep_vba=False)
    ws = wb[HOJA]
    filas = list(ws.iter_rows(min_row=FILA_GRUPOS, max_row=FILA_HEADERS, values_only=True))
    grupos, headers = filas[0], filas[1]
    wb.close()

    # propagar el nombre de grupo hacia la derecha
    banda_actual = None
    catalogo: list[dict] = []
    for i, (g, h) in enumerate(zip(grupos, headers), start=1):
        if g:
            banda_actual = str(g).strip()
        if not h:
            continue
        catalogo.append(
            {
                "col": get_column_letter(i),
                "ordinal": i,
                "banda": banda_actual,
                "header": str(h).strip(),
                "rol": _rol(str(h)),
            }
        )
    return catalogo


def agrupar_tripletas(catalogo: list[dict]) -> list[dict]:
    """Agrupa columnas consecutivas en tripletas (meta, resultado, cumplimiento)."""
    tripletas, buffer = [], {}
    for c in catalogo:
        buffer[c["rol"]] = c
        if c["rol"] == "cumplimiento" and buffer:
            tripletas.append(
                {
                    "banda": c["banda"],
                    "indicador": buffer.get("resultado", c)["header"],
                    "meta_col": buffer.get("meta", {}).get("col"),
                    "resultado_col": buffer.get("resultado", {}).get("col"),
                    "cumplimiento_col": c["col"],
                }
            )
            buffer = {}
    return tripletas


def imprimir_catalogo(ruta_xlsm: Path) -> None:
    for c in leer_catalogo(ruta_xlsm):
        print(f"{c['col']:>4} ({c['ordinal']:>3})  [{c['banda']}]  {c['rol']:<12}  {c['header'][:80]}")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--matriz", required=True)
    p.add_argument("--yaml", action="store_true", help="emite YAML de tripletas por stdout")
    args = p.parse_args(argv)

    ruta = Path(args.matriz)
    if args.yaml:
        import yaml

        tripletas = agrupar_tripletas(leer_catalogo(ruta))
        print(yaml.safe_dump({"indicadores_generado": tripletas}, allow_unicode=True, sort_keys=False))
    else:
        imprimir_catalogo(ruta)
    return 0


if __name__ == "__main__":
    sys.exit(main())
